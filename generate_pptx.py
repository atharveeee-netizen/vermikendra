import os
from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
# Set widescreen 16:9 dimensions (13.333 x 7.5 inches)
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6] # Blank layout

slide_images = [
    'slide_1_cover_complete_1789389305410.jpg',
    'slide_2_image2_style_master.jpg',
    'slide_3_research_journal_definitive.jpg',
    'slide_4_impact_research_journal.jpg',
    'slide_6_feasibility_risk_master.jpg',
    'slide_5_team_research_journal.jpg',
    'slide_7_references_master.jpg'
]

for img_path in slide_images:
    slide = prs.slides.add_slide(blank_slide_layout)
    # Add picture covering 100% of the slide canvas
    slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

pptx_out = 'Vermikendra_Complete_7_Slides_Presentation.pptx'
prs.save(pptx_out)
print(f'Generated complete PowerPoint presentation: {pptx_out}')
